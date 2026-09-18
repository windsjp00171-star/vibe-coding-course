"""把講師私人題庫 teacher/kahoot-bank.js 匯出成 Kahoot 可匯入的 .xlsx（每組一個檔）。

執行：python scripts/export_kahoot.py
輸出：teacher/kahoot/*.xlsx（私人 repo；含答案，不可放公開網站）

版面仿照 Kahoot 官方匯入範本：第 8 列是欄位標題、第 9 列開始是題目，
B～H 欄依序為 題目／選項 1～4／時限／正確答案。
若 Kahoot 改版導致匯入失敗：到 Kahoot 建立測驗 → Import spreadsheet 下載官方範本，
把本檔 B～H 欄的題目整塊複製貼到範本同樣位置即可。
"""
import hashlib
import json
import pathlib
import random
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parent.parent
BANK = ROOT / "teacher" / "kahoot-bank.js"
OUT_DIR = ROOT / "teacher" / "kahoot"

MAX_QUESTION = 95   # 各家說明不一（95／120），取較嚴格者
MAX_ANSWER = 60     # 各家說明不一（60／75），取較嚴格者
TIME_LIMITS = {5, 10, 20, 30, 60, 90, 120, 240}
DEFAULT_TIME = 20
HEADER_ROW = 8
HEADERS = [
    "Question - max 120 characters",
    "Answer 1 - max 75 characters",
    "Answer 2 - max 75 characters",
    "Answer 3 - max 75 characters",
    "Answer 4 - max 75 characters",
    "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs",
    "Correct answer(s) - choose at least one",
]


def load_bank() -> dict:
    if not BANK.exists():
        sys.exit("找不到 teacher/kahoot-bank.js：請先把私人講師 repo clone 到 teacher/ 資料夾")
    out = subprocess.run(
        ["node", "-e", "process.stdout.write(JSON.stringify(require(process.argv[1])))", str(BANK)],
        capture_output=True, text=True, encoding="utf-8", check=True,
    ).stdout
    return json.loads(out)


def validate(sets: dict) -> list[str]:
    problems = []
    for key, s in sets.items():
        for i, q in enumerate(s["questions"], 1):
            where = f"{key} 第 {i} 題"
            if len(q["q"]) > MAX_QUESTION:
                problems.append(f"{where}：題目 {len(q['q'])} 字，超過 {MAX_QUESTION}")
            if not 2 <= len(q["a"]) <= 4:
                problems.append(f"{where}：選項要 2～4 個")
            for a in q["a"]:
                if len(a) > MAX_ANSWER:
                    problems.append(f"{where}：選項「{a}」{len(a)} 字，超過 {MAX_ANSWER}")
            if not 0 <= q["ok"] < len(q["a"]):
                problems.append(f"{where}：正確答案索引超出範圍")
            if q.get("t", DEFAULT_TIME) not in TIME_LIMITS:
                problems.append(f"{where}：時限只能是 {sorted(TIME_LIMITS)}")
    return problems


def shuffled(q: dict) -> tuple[list[str], int]:
    """依題目文字決定固定的亂序：重新匯出時順序不變，正確答案也不會總在同一格。"""
    seed = int(hashlib.sha256(q["q"].encode("utf-8")).hexdigest(), 16)
    order = list(range(len(q["a"])))
    random.Random(seed).shuffle(order)
    return [q["a"][i] for i in order], order.index(q["ok"])


def write_set(key: str, s: dict) -> pathlib.Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = f"Vibe Coding 實戰課．{s['title']}"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B4"] = "匯入方式：Kahoot 建立測驗 → Import spreadsheet → 上傳本檔。若失敗，改用官方範本並貼上 B～H 欄。"
    for col, title in enumerate(HEADERS, start=2):
        cell = ws.cell(row=HEADER_ROW, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for n, q in enumerate(s["questions"], start=1):
        row = HEADER_ROW + n
        answers, correct = shuffled(q)
        ws.cell(row=row, column=1, value=n)
        ws.cell(row=row, column=2, value=q["q"])
        for i, a in enumerate(answers):
            ws.cell(row=row, column=3 + i, value=a)
        ws.cell(row=row, column=7, value=q.get("t", DEFAULT_TIME))
        ws.cell(row=row, column=8, value=correct + 1)
    for col, width in zip("ABCDEFGH", (5, 48, 26, 26, 26, 26, 14, 14)):
        ws.column_dimensions[col].width = width
    path = OUT_DIR / f"kahoot-{key}.xlsx"
    wb.save(path)
    return path


def main() -> None:
    sets = load_bank()
    problems = validate(sets)
    if problems:
        print("題庫有問題，沒有產生任何檔案：")
        print("\n".join(f"  - {p}" for p in problems))
        sys.exit(1)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    positions = []
    for key, s in sets.items():
        write_set(key, s)
        positions += [shuffled(q)[1] + 1 for q in s["questions"]]
    total = sum(len(s["questions"]) for s in sets.values())
    spread = {p: positions.count(p) for p in sorted(set(positions))}
    print(f"已匯出 {len(sets)} 組、{total} 題到 teacher/kahoot/；正確答案位置分布：{spread}")


if __name__ == "__main__":
    main()
